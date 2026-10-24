"""訴願書範本產生（民眾版）。[純函式，不呼叫 LLM / 不連 AWS]

依 citizen.check_case() 的結果與使用者填的資料，產生固定格式的訴願書範本。

設計原則（見 citizen-checker-plan.md §5）：
  - **固定格式**：不讓模型自由發揮，段落結構寫死，確保完整、可預期。
  - **不編造內容**：使用者沒填的欄位一律用「【請填寫：○○】」佔位。
  - **findings 只提供方向**：在理由段落列出可參考主張的方向，一律標註
    「請自行確認是否適用」，不保證成立。
  - 輸出 Markdown 字串；另提供 to_docx() 產生 .docx（需 python-docx，選用）。
"""

from __future__ import annotations

import datetime as _dt

from . import timeline as _tl


def _ph(label: str) -> str:
    """產生佔位符。"""
    return f"【請填寫：{label}】"


def _val(form: dict, key: str, label: str) -> str:
    """取表單值；空值回傳佔位符，絕不自動編造。"""
    v = form.get(key)
    if v is None:
        return _ph(label)
    s = str(v).strip()
    return s if s else _ph(label)


def _today_roc(today: _dt.date | None = None) -> str:
    return _tl.format_roc(today or _dt.date.today())


# 訴願書必備段落標題（供測試與 UI 驗證完整性）
SECTION_TITLES = (
    "訴願人",
    "原行政處分機關",
    "收受原處分日期",
    "訴願請求事項",
    "事實及理由",
    "證據",
    "受理訴願機關",
    "中華民國",            # 具文日期列（以「中華民國○年○月○日」開頭）
)


def build_markdown(form: dict, result: dict, today: _dt.date | None = None) -> str:
    """產生訴願書 Markdown 字串。

    form：民眾表單（與 citizen.check_case 同一份）。
    result：citizen.check_case 的回傳。
    """
    form = form or {}
    result = result or {}

    petitioner = _val(form, "petitioner_name", "訴願人姓名")
    petitioner_id = _val(form, "petitioner_id", "身分證字號／統一編號")
    petitioner_addr = _val(form, "petitioner_address", "訴願人地址（送達代收地址）")
    petitioner_phone = _val(form, "petitioner_phone", "聯絡電話")

    agency = _val(form, "agency", "原處分機關")
    doc_no = _val(form, "doc_no", "原處分書字號")
    disposition_date = _fmt_date_or_ph(form.get("disposition_date"), "原處分作成日期")
    service_date = _fmt_date_or_ph(form.get("service_date"), "收受原處分日期")

    deadline = (result.get("deadline") or {}).get("deadline_roc")
    deadline_line = deadline if deadline else _ph("訴願期限（請依處分書計算）")

    lines: list[str] = []
    lines.append("# 訴願書")
    lines.append("")
    lines.append("> 本範本由自助檢查工具產生，僅供參考，不構成法律意見。")
    lines.append("> 請於寄出前逐項確認內容正確，並刪除所有【請填寫】佔位說明。")
    lines.append("")

    # 一、訴願人
    lines.append("## 訴願人")
    lines.append("")
    lines.append(f"- 姓名／名稱：{petitioner}")
    lines.append(f"- 身分證字號／統一編號：{petitioner_id}")
    lines.append(f"- 地址：{petitioner_addr}")
    lines.append(f"- 聯絡電話：{petitioner_phone}")
    lines.append("")

    # 二、原行政處分機關
    lines.append("## 原行政處分機關")
    lines.append("")
    lines.append(f"- 機關：{agency}")
    lines.append(f"- 處分書字號：{doc_no}")
    lines.append(f"- 處分作成日期：{disposition_date}")
    lines.append("")

    # 三、收受原處分日期
    lines.append("## 收受原處分日期")
    lines.append("")
    lines.append(f"- 收受（送達）日期：{service_date}")
    lines.append(f"- 本件訴願期限（試算）：{deadline_line}")
    lines.append("  （期限以受理機關實際收到訴願書之日為準，非郵戳日；請提早以掛號寄出。）")
    lines.append("")

    # 四、訴願請求事項
    lines.append("## 訴願請求事項")
    lines.append("")
    requests = form.get("requests")
    if isinstance(requests, str) and requests.strip():
        lines.append(requests.strip())
    elif isinstance(requests, list) and requests:
        for r in requests:
            lines.append(f"1. {r}")
    else:
        lines.append(f"1. 請求撤銷{agency}所為之原處分（{doc_no}）。")
        lines.append(f"2. {_ph('其他具體請求，例如：請求變更處分、發回原機關另為適法處分等')}")
    lines.append("")

    # 五、事實及理由
    lines.append("## 事實及理由")
    lines.append("")
    lines.append("### 一、事實經過")
    lines.append("")
    facts = form.get("facts")
    if isinstance(facts, str) and facts.strip():
        lines.append(facts.strip())
    else:
        lines.append(_ph("請簡述收到處分的經過，例如：何時、因何事由受何處分"))
    lines.append("")

    lines.append("### 二、不服之理由")
    lines.append("")
    reasons = form.get("reasons")
    if isinstance(reasons, str) and reasons.strip():
        lines.append(reasons.strip())
    else:
        lines.append(_ph("請說明您為何不服原處分（事實認定、法律適用或裁量等）"))
    lines.append("")

    # findings → 可參考主張方向（§5.3）
    findings = result.get("findings") or []
    if findings:
        lines.append("### 三、可參考之主張方向（請自行確認是否適用）")
        lines.append("")
        lines.append("> 以下為自助檢查工具就處分書客觀資訊提出的提醒，**非法律意見**，")
        lines.append("> 是否適用請自行確認或洽詢專業人士後，再決定是否納入主張。")
        lines.append("")
        for f in findings:
            title = f.get("title", "")
            plain = f.get("plain", "")
            basis = f.get("basis", "")
            item = f"- **{title}**：{plain}"
            if basis:
                item += f"（參考依據：{basis}）"
            lines.append(item)
        lines.append("")

    # 六、證據
    lines.append("## 證據")
    lines.append("")
    evidence = form.get("evidence")
    if isinstance(evidence, str) and evidence.strip():
        lines.append(evidence.strip())
    elif isinstance(evidence, list) and evidence:
        for e in evidence:
            lines.append(f"1. {e}")
    else:
        lines.append("1. 原處分書影本一份。")
        lines.append(f"2. {_ph('其他證據，例如：送達信封、相關證明文件')}")
    lines.append("")

    # 七、受理訴願機關
    lines.append("## 受理訴願機關")
    lines.append("")
    receiving = form.get("receiving_agency")
    if isinstance(receiving, str) and receiving.strip():
        lines.append(f"- {receiving.strip()}")
    else:
        lines.append(f"- {_ph('受理訴願機關（通常為原處分機關之上級機關，請以處分書教示為準）')}")
    lines.append("- 提出方式：建議經由原處分機關或逕送受理訴願機關，並以掛號寄出、保留收據。")
    lines.append("")

    # 具文日期
    lines.append("---")
    lines.append("")
    lines.append(f"此致")
    lines.append("")
    lines.append(f"具文人（訴願人）：{petitioner}（簽名或蓋章）")
    lines.append("")
    lines.append(f"中華民國{_roc_ymd(today)}")
    lines.append("")

    return "\n".join(lines)


def _roc_ymd(today: _dt.date | None) -> str:
    d = today or _dt.date.today()
    return f"{d.year - 1911}年{d.month}月{d.day}日"


def _fmt_date_or_ph(value, label: str) -> str:
    """把日期值格式化成民國年；空值或無法解析回佔位符。"""
    if value is None or str(value).strip() == "":
        return _ph(label)
    d = _tl.parse_date(value)
    if d is None:
        # 有填但格式無法辨識：保留原文並提醒
        return f"{str(value).strip()}（請確認日期格式）"
    return _tl.format_roc(d)


# ---------------------------------------------------------------------------
# .docx 輸出（選用，需 python-docx）
# ---------------------------------------------------------------------------
def to_docx(form: dict, result: dict, path: str, today: _dt.date | None = None) -> str:
    """把訴願書寫成 .docx 檔。需安裝 python-docx，未安裝則丟 RuntimeError。

    回傳寫出的檔案路徑。內容與 build_markdown 相同，但轉為 Word 段落/標題樣式。
    """
    try:
        from docx import Document
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError(
            "產生 .docx 需安裝 python-docx：pip install python-docx"
        ) from exc

    md = build_markdown(form, result, today=today)
    doc = Document()
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.style = doc.styles["Intense Quote"] if "Intense Quote" in [s.name for s in doc.styles] else p.style
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line == "---":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
    doc.save(path)
    return path
