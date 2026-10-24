"""民眾版「收到處分怎麼辦」自助檢查工具（手機優先單頁）。

與承辦人版（app/streamlit_app.py）完全分離：
  - 只呼叫 core.citizen / core.citizen_template 的純規則，**不呼叫 LLM、不連 AWS、不需 boto3**。
  - 不寫檔、不記錄任何使用者輸入（僅用 st.session_state 暫存於瀏覽器工作階段）。

流程三步：① 填處分書資料 → ② 看結果 → ③ 下載訴願書範本。

執行：
    python -m streamlit run app/citizen_app.py
"""

from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import streamlit as st  # noqa: E402

from core import citizen  # noqa: E402
from core import citizen_template  # noqa: E402


DISCLAIMER = (
    "本工具僅依您輸入的資料做期限試算與格式檢查，**不構成法律意見**。"
    "期限攸關權益，請以處分書記載及受理機關認定為準，並盡早寄出；"
    "如有疑問，可洽各地方政府法律扶助或法律扶助基金會。"
)

URGENCY_STYLE = {
    "overdue": ("#b00020", "可能已逾期"),
    "urgent":  ("#e65100", "非常緊急"),
    "soon":    ("#e6a100", "期限將至"),
    "ok":      ("#2e7d32", "尚在期限內"),
    "unknown": ("#555555", "資料不足，無法試算"),
}


# ============ 頁面設定（手機優先：單欄、窄版）============
st.set_page_config(
    page_title="收到處分怎麼辦？訴願期限自助檢查",
    page_icon="📮",
    layout="centered",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
    <style>
      .block-container {max-width: 720px; padding-top: 1.2rem;}
      .disclaimer {background:#fff8e1; border:1px solid #ffe082; border-radius:10px;
                   padding:12px 14px; font-size:0.9rem; color:#5d4037; line-height:1.6;}
      .bigbanner {border-radius:14px; padding:20px 18px; color:#fff; text-align:center;
                  margin:8px 0 16px;}
      .bigbanner .big {font-size:2.2rem; font-weight:800; line-height:1.2;}
      .bigbanner .sub {font-size:1rem; opacity:0.95; margin-top:6px;}
      .hint {color:#607d8b; font-size:0.82rem;}
      .stepnav {text-align:center; color:#90a4ae; font-size:0.85rem; margin-bottom:4px;}
    </style>
    """,
    unsafe_allow_html=True,
)


def _disclaimer_block():
    st.markdown(f'<div class="disclaimer">⚠️ {DISCLAIMER}</div>', unsafe_allow_html=True)


# ============ Session 狀態 ============
if "step" not in st.session_state:
    st.session_state.step = 1
if "form" not in st.session_state:
    st.session_state.form = {}
if "result" not in st.session_state:
    st.session_state.result = None


def _goto(step: int):
    st.session_state.step = step


# ============ 頁首 ============
st.title("📮 收到處分怎麼辦？")
st.caption("訴願期限自助試算 ＋ 處分書檢查 ＋ 訴願書範本。免費、不儲存您的任何資料。")
_disclaimer_block()
st.markdown(f'<div class="stepnav">步驟 {st.session_state.step} / 3</div>',
            unsafe_allow_html=True)


# ============ 步驟一：填資料 ============
def render_step1():
    st.subheader("① 填入處分書上的資料")
    st.markdown('<span class="hint">欄位都可留空，填越多、檢查越完整。</span>',
                unsafe_allow_html=True)
    f = st.session_state.form

    with st.form("case_form"):
        agency = st.text_input(
            "處分機關", value=f.get("agency", ""),
            help="這在處分書哪裡？通常在最上方的機關全銜，或蓋章處。")
        doc_no = st.text_input(
            "處分書字號", value=f.get("doc_no", ""),
            help="這在處分書哪裡？多在右上角，例如「○○字第…號」。")

        col1, col2 = st.columns(2)
        with col1:
            disposition_date = st.text_input(
                "處分作成日期", value=f.get("disposition_date", ""),
                placeholder="例：115/09/28",
                help="這在處分書哪裡？多在最後具文處或字號旁的日期。")
        with col2:
            service_date = st.text_input(
                "您收到（送達）的日期", value=f.get("service_date", ""),
                placeholder="例：115/10/02",
                help="這在處分書哪裡？以您實際收到的日期為準，可參考信封郵戳或招領通知。")

        service_method = st.radio(
            "您是怎麼收到的？",
            options=[citizen.SERVICE_PERSONAL, citizen.SERVICE_AGENT,
                     citizen.SERVICE_DEPOSIT, citizen.SERVICE_UNKNOWN],
            index=[citizen.SERVICE_PERSONAL, citizen.SERVICE_AGENT,
                   citizen.SERVICE_DEPOSIT, citizen.SERVICE_UNKNOWN].index(
                       f.get("service_method", citizen.SERVICE_PERSONAL)),
            help="寄存送達（郵局或派出所招領）的生效日會往後算 10 天，影響期限。")

        act_date = st.text_input(
            "違規行為日期（選填）", value=f.get("act_date", ""),
            placeholder="例：115/08/01",
            help="這在處分書哪裡？多在「事實」或「違規時間」欄位。用來檢查是否超過裁處期限。")

        legal_basis = st.text_input(
            "處分書寫的法條依據（選填）", value=f.get("legal_basis", ""),
            placeholder="例：道路交通管理處罰條例第56條",
            help="這在處分書哪裡？多在「法令依據」或理由段落。")

        col3, col4 = st.columns(2)
        with col3:
            penalty_amount = st.text_input(
                "罰鍰金額（選填）", value=str(f.get("penalty_amount", "") or ""),
                placeholder="例：900")
        with col4:
            has_remedy = st.selectbox(
                "處分書有寫「幾日內可提起訴願」嗎？",
                options=["不確定", "有", "沒有"],
                index=["不確定", "有", "沒有"].index(f.get("_has_remedy_label", "不確定")),
                help="就是救濟教示。沒寫的話，提起期間可能延長，但仍建議以 30 天為準。")

        already_label = st.selectbox(
            "您已經寄出訴願書了嗎？",
            options=["還沒", "已寄出"],
            index=["還沒", "已寄出"].index(f.get("_already_label", "還沒")))
        already_date = ""
        if already_label == "已寄出":
            already_date = st.text_input(
                "寄出日期", value=f.get("_already_date", ""), placeholder="例：115/10/20")

        submitted = st.form_submit_button("開始檢查 →", use_container_width=True)

    if submitted:
        has_remedy_notice = {"有": True, "沒有": False, "不確定": None}[has_remedy]
        form = {
            "agency": agency, "doc_no": doc_no,
            "disposition_date": disposition_date, "service_date": service_date,
            "service_method": service_method, "act_date": act_date,
            "legal_basis": legal_basis,
            "penalty_amount": penalty_amount,
            "has_remedy_notice": has_remedy_notice,
            # UI 記憶用（底線前綴不進 core）
            "_has_remedy_label": has_remedy,
            "_already_label": already_label,
            "_already_date": already_date,
        }
        if already_label == "已寄出" and already_date.strip():
            form["already_filed"] = {"date": already_date.strip()}
        st.session_state.form = form
        st.session_state.result = citizen.check_case(form, today=_dt.date.today())
        _goto(2)
        st.rerun()


# ============ 步驟二：看結果 ============
def render_step2():
    result = st.session_state.result
    if not result:
        _goto(1)
        st.rerun()
        return

    st.subheader("② 檢查結果")

    urgency = result.get("urgency", "unknown")
    days_left = result.get("days_left")
    deadline = result.get("deadline") or {}
    color, label = URGENCY_STYLE.get(urgency, URGENCY_STYLE["unknown"])

    # 大字期限橫幅
    if urgency == "unknown":
        big, sub = "資料不足", "請補齊「收到日期」才能試算期限"
    elif urgency == "overdue":
        big = f"已超過 {abs(days_left)} 天"
        sub = f"試算期限：{deadline.get('deadline_roc', '—')}（是否逾期以受理機關認定為準）"
    else:
        big = f"還剩 {days_left} 天"
        sub = f"訴願期限（試算）：{deadline.get('deadline_roc', '—')}"

    st.markdown(
        f'<div class="bigbanner" style="background:{color}">'
        f'<div class="big">{big}</div><div class="sub">{label}｜{sub}</div></div>',
        unsafe_allow_html=True,
    )

    # 計算說明（可追溯）
    if deadline.get("explanation"):
        with st.expander("期限是怎麼算出來的？", expanded=False):
            st.write(deadline["explanation"])
            if deadline.get("basis"):
                st.caption(f"依據：{deadline['basis']}")

    # 無救濟教示的一年延長說明
    ext = deadline.get("extended_one_year")
    if ext:
        st.info(
            f"處分書若未記載救濟教示，提起期間可能延長至 **{ext.get('deadline_roc', '—')}**"
            f"（{ext.get('basis', '')}）。\n\n"
            "但認定存在風險，**強烈建議仍以 30 天為準、儘早寄出**。"
        )

    # 警告
    for w in result.get("warnings", []):
        st.warning(w)

    # 瑕疵發現（白話）
    findings = result.get("findings", [])
    if findings:
        st.markdown("#### 處分書可能有的瑕疵（僅供參考）")
        st.caption("以下為依客觀資訊提出的提醒，一律用「可能」，是否適用請自行確認。")
        for fnd in findings:
            with st.container(border=True):
                st.markdown(f"**{fnd.get('title', '')}**")
                st.write(fnd.get("plain", ""))
                if fnd.get("basis"):
                    st.caption(f"參考依據：{fnd['basis']}")

    # 下一步
    st.markdown("#### 下一步建議")
    for s in result.get("next_steps", []):
        st.markdown(f"- {s}")

    col1, col2 = st.columns(2)
    with col1:
        if st.button("← 回去修改", use_container_width=True):
            _goto(1)
            st.rerun()
    with col2:
        if st.button("產生訴願書範本 →", use_container_width=True):
            _goto(3)
            st.rerun()

    _disclaimer_block()


# ============ 步驟三：下載訴願書範本 ============
def render_step3():
    result = st.session_state.result
    form = st.session_state.form
    if not result:
        _goto(1)
        st.rerun()
        return

    st.subheader("③ 訴願書範本")
    st.caption("以下為固定格式範本，沒填的欄位會標示【請填寫：…】，請務必自行補齊與確認。")

    # 補充訴願人資料（這些不影響期限計算，只用於產生範本）
    with st.expander("補充訴願人資料（選填，讓範本更完整）", expanded=False):
        form["petitioner_name"] = st.text_input(
            "訴願人姓名／名稱", value=form.get("petitioner_name", ""))
        form["petitioner_id"] = st.text_input(
            "身分證字號／統一編號", value=form.get("petitioner_id", ""))
        form["petitioner_address"] = st.text_input(
            "地址（送達代收地址）", value=form.get("petitioner_address", ""))
        form["petitioner_phone"] = st.text_input(
            "聯絡電話", value=form.get("petitioner_phone", ""))
        form["receiving_agency"] = st.text_input(
            "受理訴願機關（選填）", value=form.get("receiving_agency", ""))
    st.session_state.form = form

    md = citizen_template.build_markdown(form, result, today=_dt.date.today())

    st.markdown("##### 預覽")
    with st.container(border=True):
        st.markdown(md)

    st.download_button(
        "下載 Markdown（.md）", data=md.encode("utf-8"),
        file_name="訴願書範本.md", mime="text/markdown",
        use_container_width=True,
    )

    # .docx（python-docx 可用時才提供）
    try:
        import io
        from docx import Document  # noqa: F401
        # 寫到記憶體 buffer，不落地到磁碟
        buf = _docx_buffer(form, result)
        if buf is not None:
            st.download_button(
                "下載 Word（.docx）", data=buf,
                file_name="訴願書範本.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
    except ImportError:
        st.caption("（此環境未安裝 python-docx，僅提供 Markdown 下載。）")

    if st.button("← 回結果", use_container_width=True):
        _goto(2)
        st.rerun()

    _disclaimer_block()


def _docx_buffer(form: dict, result: dict):
    """把訴願書 .docx 寫進記憶體 buffer（不寫磁碟，符合『不落地』原則）。"""
    import io
    from docx import Document

    md = citizen_template.build_markdown(form, result, today=_dt.date.today())
    doc = Document()
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line == "---":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ============ 路由 ============
_STEP = st.session_state.step
if _STEP == 1:
    render_step1()
elif _STEP == 2:
    render_step2()
else:
    render_step3()
