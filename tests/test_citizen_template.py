"""core/citizen_template.py 訴願書範本產生測試（任務 5.5）。

涵蓋：必要段落齊全、沒填的欄位都有佔位符、findings 轉成可參考方向、
.docx 輸出（未安裝 python-docx 時略過）、不自動編造內容。
"""
from __future__ import annotations

import datetime as dt
import importlib.util

import pytest

from core import citizen
from core import citizen_template as ct


def _result_full():
    form = {
        "agency": "新北市政府交通事件裁決處",
        "doc_no": "北監裁字第1140000001號",
        "disposition_date": "2026-09-28",
        "service_date": "2026-10-02",
        "service_method": citizen.SERVICE_PERSONAL,
        "has_remedy_notice": True,
        "legal_basis": "",          # 觸發 D1 finding
    }
    return form, citizen.check_case(form, today=dt.date(2026, 10, 2))


def test_markdown_contains_all_required_sections():
    form, result = _result_full()
    md = ct.build_markdown(form, result, today=dt.date(2026, 10, 2))
    for title in ct.SECTION_TITLES:
        assert title in md, f"缺少段落：{title}"


def test_missing_fields_use_placeholders():
    # 完全空白的表單：訴願人、受理機關等必有佔位符
    md = ct.build_markdown({}, citizen.check_case({}, today=dt.date(2026, 10, 2)),
                           today=dt.date(2026, 10, 2))
    assert "【請填寫：訴願人姓名】" in md
    assert "【請填寫：原處分機關】" in md
    assert "【請填寫：身分證字號／統一編號】" in md


def test_no_fabrication_when_fields_empty():
    # 沒填理由/事實時，不應出現任何捏造內容，只有佔位符
    md = ct.build_markdown({}, {}, today=dt.date(2026, 10, 2))
    assert _ph_count(md) >= 5          # 多個欄位都該是佔位符
    # 不應出現像是自動編造的具體事實
    assert "【請填寫：請簡述收到處分的經過" in md


def test_findings_rendered_as_reference_directions():
    form, result = _result_full()
    md = ct.build_markdown(form, result, today=dt.date(2026, 10, 2))
    assert "可參考之主張方向" in md
    assert "請自行確認是否適用" in md
    # D1 的白話標題應出現
    assert "處分書記載事項可能不完備" in md


def test_filled_fields_appear_in_output():
    form, result = _result_full()
    md = ct.build_markdown(form, result, today=dt.date(2026, 10, 2))
    assert "新北市政府交通事件裁決處" in md
    assert "北監裁字第1140000001號" in md
    assert "民國115年10月2日" in md       # 收受日期格式化為民國年
    assert "中華民國115年10月2日" in md    # 具文日期（注入 today=2026-10-02）


def test_date_roc_formatting():
    assert ct._fmt_date_or_ph("2026-10-02", "x") == "民國115年10月2日"
    assert ct._fmt_date_or_ph("", "收受日期") == "【請填寫：收受日期】"
    # 無法解析 → 保留原文並提醒
    out = ct._fmt_date_or_ph("亂寫", "收受日期")
    assert "亂寫" in out and "確認" in out


def _ph_count(md: str) -> int:
    return md.count("【請填寫：")


# ---------- .docx 輸出 ----------
@pytest.mark.skipif(
    importlib.util.find_spec("docx") is None,
    reason="未安裝 python-docx，略過 .docx 測試",
)
def test_to_docx_writes_file(tmp_path):
    form, result = _result_full()
    out = tmp_path / "appeal.docx"
    path = ct.to_docx(form, result, str(out), today=dt.date(2026, 10, 2))
    assert out.exists()
    assert out.stat().st_size > 0
    assert path == str(out)

    # 確認內容包含關鍵段落（讀回 docx 純文字）
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "訴願書" in text
    assert "新北市政府交通事件裁決處" in text
